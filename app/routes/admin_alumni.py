import io
from datetime import datetime
from flask import Blueprint, render_template, redirect, url_for, flash, request, send_file
from flask_login import login_required
from app import mysql
from app.utils.helpers import save_image, delete_image, paginate
import openpyxl

admin_alumni_bp = Blueprint('admin_alumni', __name__)


@admin_alumni_bp.route('/alumni')
@login_required
def index():
    q = request.args.get('q', '').strip()
    angkatan = request.args.get('angkatan', '').strip()
    jurusan = request.args.get('jurusan', '').strip()
    page = int(request.args.get('page', 1))

    cur = mysql.connection.cursor()

    # Build query
    conditions = []
    params = []

    if q:
        conditions.append("(nama LIKE %s OR pekerjaan LIKE %s OR perusahaan LIKE %s)")
        params += [f'%{q}%', f'%{q}%', f'%{q}%']
    if angkatan:
        conditions.append("angkatan = %s")
        params.append(angkatan)
    if jurusan:
        conditions.append("jurusan = %s")
        params.append(jurusan)

    where = "WHERE " + " AND ".join(conditions) if conditions else ""
    cur.execute(f"SELECT * FROM alumni {where} ORDER BY created_at DESC", params)
    all_alumni = cur.fetchall()

    # Get filter options
    cur.execute("SELECT DISTINCT angkatan FROM alumni ORDER BY angkatan DESC")
    angkatan_list = [r['angkatan'] for r in cur.fetchall()]

    cur.execute("SELECT DISTINCT jurusan FROM alumni ORDER BY jurusan")
    jurusan_list = [r['jurusan'] for r in cur.fetchall()]

    cur.close()

    pagination = paginate(all_alumni, page, 15)

    return render_template('admin/alumni/index.html',
                           pagination=pagination,
                           q=q,
                           angkatan_filter=angkatan,
                           jurusan_filter=jurusan,
                           angkatan_list=angkatan_list,
                           jurusan_list=jurusan_list)


@admin_alumni_bp.route('/alumni/tambah', methods=['GET', 'POST'])
@login_required
def tambah():
    if request.method == 'POST':
        nama = request.form.get('nama', '').strip()
        angkatan = request.form.get('angkatan', '').strip()
        jurusan = request.form.get('jurusan', '').strip()

        if not nama or not angkatan or not jurusan:
            flash('Nama, angkatan, dan jurusan wajib diisi.', 'danger')
            return render_template('admin/alumni/form.html', alumni=None)

        if not angkatan.isdigit() or not (1950 <= int(angkatan) <= 2100):
            flash('Angkatan harus berupa tahun antara 1950-2100.', 'danger')
            return render_template('admin/alumni/form.html', alumni=None)

        pekerjaan = request.form.get('pekerjaan', '').strip()
        perusahaan = request.form.get('perusahaan', '').strip()
        kota_domisili = request.form.get('kota_domisili', '').strip()
        email = request.form.get('email', '').strip()
        linkedin = request.form.get('linkedin', '').strip()
        bio = request.form.get('bio', '').strip()
        is_unggulan = 1 if request.form.get('is_unggulan') else 0

        foto_filename = None
        if 'foto' in request.files and request.files['foto'].filename:
            foto_filename = save_image(request.files['foto'], 'alumni', (800, 800))
            if not foto_filename:
                flash('Format file foto tidak valid. Gunakan png, jpg, jpeg, gif, atau webp.', 'danger')
                return render_template('admin/alumni/form.html', alumni=None)

        cur = mysql.connection.cursor()
        cur.execute("""
            INSERT INTO alumni (nama, angkatan, jurusan, foto, pekerjaan, perusahaan,
                                kota_domisili, email, linkedin, bio, is_unggulan)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (nama, angkatan, jurusan, foto_filename, pekerjaan, perusahaan,
              kota_domisili, email, linkedin, bio, is_unggulan))
        mysql.connection.commit()
        cur.close()

        flash('Data alumni berhasil ditambahkan.', 'success')
        return redirect(url_for('admin_alumni.index'))

    return render_template('admin/alumni/form.html', alumni=None)


@admin_alumni_bp.route('/alumni/edit/<int:id>', methods=['GET', 'POST'])
@login_required
def edit(id):
    cur = mysql.connection.cursor()
    cur.execute("SELECT * FROM alumni WHERE id = %s", (id,))
    alumni = cur.fetchone()
    cur.close()

    if not alumni:
        flash('Data alumni tidak ditemukan.', 'danger')
        return redirect(url_for('admin_alumni.index'))

    if request.method == 'POST':
        nama = request.form.get('nama', '').strip()
        angkatan = request.form.get('angkatan', '').strip()
        jurusan = request.form.get('jurusan', '').strip()

        if not nama or not angkatan or not jurusan:
            flash('Nama, angkatan, dan jurusan wajib diisi.', 'danger')
            return render_template('admin/alumni/form.html', alumni=alumni)

        if not angkatan.isdigit() or not (1950 <= int(angkatan) <= 2100):
            flash('Angkatan harus berupa tahun antara 1950-2100.', 'danger')
            return render_template('admin/alumni/form.html', alumni=alumni)

        pekerjaan = request.form.get('pekerjaan', '').strip()
        perusahaan = request.form.get('perusahaan', '').strip()
        kota_domisili = request.form.get('kota_domisili', '').strip()
        email = request.form.get('email', '').strip()
        linkedin = request.form.get('linkedin', '').strip()
        bio = request.form.get('bio', '').strip()
        is_unggulan = 1 if request.form.get('is_unggulan') else 0

        foto_filename = alumni['foto']
        if 'foto' in request.files and request.files['foto'].filename:
            new_foto = save_image(request.files['foto'], 'alumni', (800, 800))
            if not new_foto:
                flash('Format file foto tidak valid.', 'danger')
                return render_template('admin/alumni/form.html', alumni=alumni)
            # Delete old photo
            if alumni['foto']:
                delete_image(alumni['foto'], 'alumni')
            foto_filename = new_foto

        cur = mysql.connection.cursor()
        cur.execute("""
            UPDATE alumni SET nama=%s, angkatan=%s, jurusan=%s, foto=%s, pekerjaan=%s,
                perusahaan=%s, kota_domisili=%s, email=%s, linkedin=%s, bio=%s, is_unggulan=%s
            WHERE id=%s
        """, (nama, angkatan, jurusan, foto_filename, pekerjaan, perusahaan,
              kota_domisili, email, linkedin, bio, is_unggulan, id))
        mysql.connection.commit()
        cur.close()

        flash('Data alumni berhasil diperbarui.', 'success')
        return redirect(url_for('admin_alumni.index'))

    return render_template('admin/alumni/form.html', alumni=alumni)


@admin_alumni_bp.route('/alumni/hapus/<int:id>', methods=['POST'])
@login_required
def hapus(id):
    cur = mysql.connection.cursor()
    cur.execute("SELECT foto FROM alumni WHERE id = %s", (id,))
    alumni = cur.fetchone()

    if alumni:
        if alumni['foto']:
            delete_image(alumni['foto'], 'alumni')
        cur.execute("DELETE FROM alumni WHERE id = %s", (id,))
        mysql.connection.commit()
        flash('Data alumni berhasil dihapus.', 'success')
    else:
        flash('Data alumni tidak ditemukan.', 'danger')

    cur.close()
    return redirect(url_for('admin_alumni.index'))


@admin_alumni_bp.route('/alumni/export')
@login_required
def export():
    cur = mysql.connection.cursor()
    cur.execute("SELECT nama, angkatan, jurusan, pekerjaan, perusahaan, kota_domisili, email, linkedin FROM alumni ORDER BY angkatan DESC, nama")
    alumni_list = cur.fetchall()
    cur.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Alumni"

    headers = ['Nama', 'Angkatan', 'Jurusan', 'Pekerjaan', 'Perusahaan', 'Kota Domisili', 'Email', 'LinkedIn']
    ws.append(headers)

    for a in alumni_list:
        ws.append([
            a['nama'], a['angkatan'], a['jurusan'],
            a['pekerjaan'] or '', a['perusahaan'] or '',
            a['kota_domisili'] or '', a['email'] or '', a['linkedin'] or ''
        ])

    # Style header
    from openpyxl.styles import Font, PatternFill
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="1E3A5F", end_color="1E3A5F", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_len + 4

    timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    filename = f"data_alumni_{timestamp}.xlsx"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(output,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True,
                     download_name=filename)
